import sys
import docx
from pptx import Presentation

def extract_ppt(path):
    print("=== PPT TEMPLATE ===")
    try:
        prs = Presentation(path)
        for i, slide in enumerate(prs.slides):
            print(f"Slide {i+1}:")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    print(f"  - {shape.text.strip()}")
    except Exception as e:
        print(f"Error reading PPT: {e}")

def extract_docx(path):
    print("\n=== DOCX TEMPLATE ===")
    try:
        doc = docx.Document(path)
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                print(f"  {p.text.strip()}")
        print("\n  [Tables in Doc]:")
        for t in doc.tables:
            for r in t.rows:
                row_data = [c.text.replace('\n', ' ') for c in r.cells]
                print(f"  - {' | '.join(row_data)}")
    except Exception as e:
        print(f"Error reading DOCX: {e}")

extract_ppt(r"C:\Users\Alena B\Downloads\SIH2026-IDEA-Presentation-Format.pptx")
extract_docx(r"C:\Users\Alena B\Downloads\SIH Project Proposal Template.docx")
